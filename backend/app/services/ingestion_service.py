"""
Resume ingestion service.

Handles PDF, DOCX, OCR-fallback, and manual text ingestion with
character-offset source spans for downstream traceability.

Public API:
    ingest_file(file, candidate_id)  → ResumeDocument
    ingest_manual(text, candidate_id) → ResumeDocument
    get_resume(resume_id)            → ResumeDocument

Requirements: 1.1–1.12
"""

from __future__ import annotations

import io
import os
import uuid
from datetime import UTC, datetime
from typing import Any

from sqlalchemy.orm import Session

from backend.app.exceptions import (
    EmptyExtractionError,
    ExtractionFailedError,
    FileTooLargeError,
    ResourceNotFoundError,
    TextEmptyError,
    TextTooLongError,
    UnsupportedFileTypeError,
)
from backend.app.models.enums import DocumentType
from backend.app.models.resume_document import ResumeDocument

# ── Config ────────────────────────────────────────────────────────────────────

MAX_FILE_SIZE_MB: int = int(os.environ.get("MAX_FILE_SIZE_MB", "10"))
MAX_MANUAL_CHARS: int = int(os.environ.get("MAX_MANUAL_CHARS", "50000"))

_ALLOWED_CONTENT_TYPES: set[str] = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}

_EXTENSION_MAP: dict[str, DocumentType] = {
    ".pdf": DocumentType.PDF,
    ".docx": DocumentType.DOCX,
}


# ── PDF extraction ────────────────────────────────────────────────────────────


def _extract_pdf(data: bytes) -> tuple[str, list[dict[str, Any]]]:
    """
    Extract text and source spans from a PDF file.

    Falls back to OCR via pytesseract when PyMuPDF extracts zero
    non-whitespace characters (scanned/image-only PDFs).
    """
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:
        raise ExtractionFailedError("PyMuPDF (fitz) is not installed.") from exc

    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ExtractionFailedError(str(exc)) from exc

    full_text = ""
    spans: list[dict[str, Any]] = []
    offset = 0

    for page in doc:
        page_text = page.get_text("text")
        if page_text:
            start = offset
            full_text += page_text
            offset += len(page_text)
            spans.append({"start": start, "end": offset, "text": page_text})

    doc.close()

    # OCR fallback: if PyMuPDF returned zero non-whitespace chars
    if not full_text.strip():
        full_text, spans = _extract_pdf_ocr(data)

    return full_text, spans


def _extract_pdf_ocr(data: bytes) -> tuple[str, list[dict[str, Any]]]:
    """OCR fallback using pytesseract for scanned PDFs."""
    try:
        import fitz  # PyMuPDF
        import pytesseract
        from PIL import Image
    except ImportError as exc:
        raise ExtractionFailedError(
            "pytesseract, Pillow, or PyMuPDF is not installed for OCR."
        ) from exc

    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ExtractionFailedError(str(exc)) from exc

    full_text = ""
    spans: list[dict[str, Any]] = []
    offset = 0

    for page in doc:
        pix = page.get_pixmap(dpi=300)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        page_text: str = pytesseract.image_to_string(img)
        if page_text:
            start = offset
            full_text += page_text
            offset += len(page_text)
            spans.append({"start": start, "end": offset, "text": page_text})

    doc.close()
    return full_text, spans


# ── DOCX extraction ───────────────────────────────────────────────────────────


def _extract_docx(data: bytes) -> tuple[str, list[dict[str, Any]]]:
    """Extract text and source spans from a DOCX file."""
    try:
        from docx import Document  # python-docx
    except ImportError as exc:
        raise ExtractionFailedError("python-docx is not installed.") from exc

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionFailedError(str(exc)) from exc

    full_text = ""
    spans: list[dict[str, Any]] = []
    offset = 0

    for para in doc.paragraphs:
        text = para.text
        if text:
            start = offset
            full_text += text + "\n"
            offset += len(text) + 1  # +1 for newline
            spans.append({"start": start, "end": offset - 1, "text": text})

    return full_text, spans


# ── Public API ────────────────────────────────────────────────────────────────


def ingest_file(
    *,
    db: Session,
    file_data: bytes,
    filename: str,
    content_type: str,
    candidate_id: uuid.UUID,
) -> ResumeDocument:
    """
    Ingest a PDF or DOCX file upload.

    Validates file size and type before attempting extraction.
    Raises appropriate AppError subclasses on any failure path,
    ensuring no partial ``ResumeDocument`` is persisted.

    Args:
        db: Active database session.
        file_data: Raw file bytes.
        filename: Original filename.
        content_type: MIME type of the upload.
        candidate_id: UUID of the owning candidate.

    Returns:
        The persisted ResumeDocument.

    Raises:
        FileTooLargeError: File exceeds MAX_FILE_SIZE_MB.
        UnsupportedFileTypeError: Not a PDF or DOCX.
        EmptyExtractionError: Extraction yielded no usable text.
        ExtractionFailedError: Library threw an unexpected error.
    """
    # ── Size check ────────────────────────────────────────────────────────
    max_bytes = MAX_FILE_SIZE_MB * 1024 * 1024
    if len(file_data) > max_bytes:
        raise FileTooLargeError(MAX_FILE_SIZE_MB)

    # ── MIME / extension check ────────────────────────────────────────────
    ext = _get_extension(filename)
    doc_type = _EXTENSION_MAP.get(ext)
    if doc_type is None and content_type not in _ALLOWED_CONTENT_TYPES:
        raise UnsupportedFileTypeError()
    if doc_type is None:
        doc_type = (
            DocumentType.PDF if content_type == "application/pdf" else DocumentType.DOCX
        )

    # ── Extract ───────────────────────────────────────────────────────────
    if doc_type == DocumentType.PDF:
        extracted_text, source_spans = _extract_pdf(file_data)
    else:
        extracted_text, source_spans = _extract_docx(file_data)

    # ── Empty check ───────────────────────────────────────────────────────
    if not extracted_text.strip():
        raise EmptyExtractionError()

    # ── Persist ───────────────────────────────────────────────────────────
    resume_doc = ResumeDocument(
        candidate_id=candidate_id,
        filename=filename,
        document_type=doc_type,
        extracted_text=extracted_text,
        source_spans=source_spans,
        created_at=datetime.now(UTC),
    )
    db.add(resume_doc)
    db.commit()
    db.refresh(resume_doc)
    return resume_doc


def ingest_manual(
    *,
    db: Session,
    text: str,
    candidate_id: uuid.UUID,
) -> ResumeDocument:
    """
    Ingest manually entered resume text.

    Accepts 1–50,000 characters. Returns 422 on empty or over-limit text.

    Args:
        db: Active database session.
        text: The candidate's resume text.
        candidate_id: UUID of the owning candidate.

    Returns:
        The persisted ResumeDocument.

    Raises:
        TextEmptyError: Text is empty or whitespace-only.
        TextTooLongError: Text exceeds 50,000 characters.
    """
    if not text or not text.strip():
        raise TextEmptyError("Resume text")

    if len(text) > MAX_MANUAL_CHARS:
        raise TextTooLongError("Resume text", MAX_MANUAL_CHARS)

    source_spans: list[dict[str, Any]] = [
        {"start": 0, "end": len(text), "text": text}
    ]

    resume_doc = ResumeDocument(
        candidate_id=candidate_id,
        filename=None,
        document_type=DocumentType.MANUAL,
        extracted_text=text,
        source_spans=source_spans,
        created_at=datetime.now(UTC),
    )
    db.add(resume_doc)
    db.commit()
    db.refresh(resume_doc)
    return resume_doc


def get_resume(*, db: Session, resume_id: uuid.UUID) -> ResumeDocument:
    """
    Retrieve a ResumeDocument by primary key.

    Raises:
        ResourceNotFoundError: No document with the given ID.
    """
    resume = db.get(ResumeDocument, resume_id)
    if resume is None:
        raise ResourceNotFoundError("ResumeDocument", str(resume_id))
    return resume


# ── Helpers ───────────────────────────────────────────────────────────────────


def _get_extension(filename: str) -> str:
    """Return the lower-cased file extension including the dot."""
    dot_idx = filename.rfind(".")
    if dot_idx == -1:
        return ""
    return filename[dot_idx:].lower()
